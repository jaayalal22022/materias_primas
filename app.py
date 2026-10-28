"""Tablero local de validación del Tomo II de la ENESEM: Materias Primas."""

from collections import Counter
from io import BytesIO
import hashlib

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import streamlit as st
from company_view import company_summary, matrix_records, prepare_batch

from engine import (
    EDIT_FIELDS, NUMERIC_EDITS, QUANTITIES, REVIEW_STATES,
    blank, conversion_for, converted_edits, decimal, display_number, evaluate,
)


st.set_page_config(page_title="ENESEM | Materias Primas", layout="wide")
st.title("Validación de Materias Primas")
st.caption("Tomo II · Revisión por registro y por empresa · Valores monetarios de solo lectura")


def load_source(content):
    workbook = load_workbook(BytesIO(content), data_only=False)
    if "MP" not in workbook or "CPC" not in workbook or "UM" not in workbook:
        raise ValueError("El archivo debe incluir las hojas MP, CPC y UM.")
    sheet = workbook["MP"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    required = {
        "id_mp", "inec_identificador_empresa", "nombre", "cpc", "unid_med",
        "ciiu4_actividad_principal", "ciiu4_actividad_secundaria", "ciiu4_tercera_actividad",
        "valor_inv_ini", "valor_impor", "valor_local", "valor_inv_fin",
        "cantidad_inv_ini", "cantidad_impor", "cantidad_local", "cantidad_inv_fin",
        "cant_terceros", "porcentaje_local",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Faltan columnas en MP: " + ", ".join(missing))
    rows, initial_edits = {}, {}
    for sheet_row in range(2, sheet.max_row + 1):
        identifier = sheet.cell(sheet_row, headers["id_mp"]).value
        if blank(identifier):
            continue
        identifier = str(identifier).strip()
        if identifier in rows:
            raise ValueError(f"El id_mp {identifier} está duplicado (fila {sheet_row}).")
        row = {name: sheet.cell(sheet_row, col).value for name, col in headers.items()}
        row["_excel_row"] = sheet_row
        rows[identifier] = row
        initial_edits[identifier] = {
            field: str(row.get(field)) if not blank(row.get(field)) else ""
            for field in EDIT_FIELDS
        }
        if initial_edits[identifier]["Estado_revision"] not in REVIEW_STATES:
            initial_edits[identifier]["Estado_revision"] = "Pendiente"
    if not rows:
        raise ValueError("La hoja MP no tiene registros con id_mp.")

    units = {}
    for code, label, target, factor, *tail in workbook["UM"].iter_rows(min_row=2, max_col=6, values_only=True):
        if code is not None and decimal(code) is not None:
            units[int(code)] = {"label": str(label or code), "target": target, "factor": factor}
    cpc = {}
    for code, description in workbook["CPC"].iter_rows(min_row=2, min_col=4, max_col=5, values_only=True):
        if not blank(code):
            cpc.setdefault(str(code).strip(), str(description or ""))
    return rows, initial_edits, units, cpc


def get_report(rows, edits, units, cpc):
    return {key: evaluate(row, edits[key], units, cpc) for key, row in rows.items()}


def export_progress(content, rows, edits, units, cpc, report):
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["MP"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    for field in EDIT_FIELDS:
        if field not in headers:
            col = sheet.max_column + 1
            sheet.cell(1, col, field)
            headers[field] = col
    for key, row in rows.items():
        for field in EDIT_FIELDS:
            cell = sheet.cell(row["_excel_row"], headers[field])
            cell.value = edits[key].get(field, "") or None
            if field == "cod_cpc_act" and cell.value is not None:
                cell.number_format = "@"
    if "Revision_app" in workbook:
        del workbook["Revision_app"]
    review = workbook.create_sheet("Revision_app", 1)
    review.append([
        "id_mp", "inec_identificador_empresa", "Materia prima", "CPC efectivo",
        "Unidad efectiva", "Estado revisión", "Alertas", "Detalle de alertas",
        "Justificación a las novedades", "Responsable",
    ])
    for key, row in rows.items():
        result = report[key]
        review.append([
            key, str(row.get("inec_identificador_empresa") or ""), str(row.get("nombre") or ""),
            result["cpc"], result["unit"], edits[key]["Estado_revision"],
            len(result["alerts"]), "\n".join(result["alerts"]),
            edits[key]["Justificación a las novedades"], edits[key]["Responsable"],
        ])
    for cell in review[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    review.freeze_panes = "D2"
    review.auto_filter.ref = review.dimensions
    for letter, width in {
        "A": 20, "B": 24, "C": 35, "D": 20, "E": 19,
        "F": 20, "G": 12, "H": 72, "I": 48, "J": 22,
    }.items():
        review.column_dimensions[letter].width = width
    for data_row in review.iter_rows(min_row=2):
        data_row[7].alignment = Alignment(wrap_text=True, vertical="top")
        data_row[8].alignment = Alignment(wrap_text=True, vertical="top")
    if workbook.calculation:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


uploaded = st.sidebar.file_uploader("Cargar matriz o archivo de avance", type=["xlsx"])
if uploaded is None:
    st.info("Carga la matriz de Materias Primas para iniciar o retomar una revisión.")
    st.stop()

content = uploaded.getvalue()
digest = hashlib.sha256(content).hexdigest()
if st.session_state.get("mp_file_digest") != digest:
    try:
        rows, initial_edits, units, cpc = load_source(content)
    except Exception as exc:
        st.error(f"No se pudo cargar la matriz: {exc}")
        st.stop()
    st.session_state.update({
        "mp_file_digest": digest, "mp_content": content, "mp_rows": rows,
        "mp_edits": initial_edits, "mp_units": units, "mp_cpc": cpc,
        "mp_revision": 0,
    })

rows = st.session_state.mp_rows
edits = st.session_state.mp_edits
units = st.session_state.mp_units
cpc = st.session_state.mp_cpc
report = get_report(rows, edits, units, cpc)

company_ids = sorted({str(row.get("inec_identificador_empresa") or "") for row in rows.values()})
company = st.sidebar.selectbox("Empresa", ["Todas"] + company_ids)
status_filter = st.sidebar.selectbox("Estado", ["Todos"] + list(REVIEW_STATES))
only_alerts = st.sidebar.checkbox("Solo registros con alertas")
search = st.sidebar.text_input("Buscar materia prima, CPC o id_mp").strip().casefold()

filtered = [key for key, row in rows.items()
            if (company == "Todas" or str(row.get("inec_identificador_empresa")) == company)
            and (status_filter == "Todos" or edits[key]["Estado_revision"] == status_filter)
            and (not only_alerts or report[key]["alerts"])
            and (not search or search in " ".join((key, str(row.get("nombre") or ""), report[key]["cpc"])).casefold())]

counts = Counter(edits[key]["Estado_revision"] for key in rows)
a, b, c, d = st.columns(4)
a.metric("Materias primas", len(rows))
b.metric("Pendientes", counts["Pendiente"])
c.metric("Validadas", counts["Validado"])
d.metric("Con alertas", sum(bool(result["alerts"]) for result in report.values()))

tab_review, tab_company = st.tabs(["Revisar registro", "Empresa por empresa"])
with tab_review:
    if not filtered:
        st.info("No hay registros para los filtros seleccionados.")
    else:
        selected = st.selectbox(
            "Registro", filtered,
            format_func=lambda key: f"{key} · {str(rows[key].get('nombre') or '')[:65]} · {report[key]['cpc']}",
        )
        row, current, result = rows[selected], edits[selected], report[selected]
        st.subheader(str(row.get("nombre") or "Materia prima"))
        st.write(f"Empresa: **{row.get('inec_identificador_empresa')}** · Registro: **{selected}**")
        with st.expander("Actividades de la empresa y Capítulo 3"):
            for kind in ("principal", "secundaria"):
                st.write(f"{kind.title()}: {row.get('ciiu4_actividad_' + kind) or '—'} · {row.get('desc_actividad_' + kind) or '—'}")
            st.write(f"Tercera: {row.get('ciiu4_tercera_actividad') or '—'} · {row.get('desc_tercera_actividad') or '—'}")

        st.caption(f"CPC original: {row.get('cpc')} · {cpc.get(str(row.get('cpc')), 'Sin descripción en catálogo')}")
        source_unit = decimal(row.get("unid_med"))
        source_code = int(source_unit) if source_unit is not None and source_unit == int(source_unit) else None
        source_unit_name = units.get(source_code, {}).get("label", "Sin unidad")
        effective_unit_name = units.get(result["unit"], {}).get("label", "Sin unidad")
        st.write(f"Unidad: **{source_unit_name}** → **{effective_unit_name}** · CPC efectivo: **{result['cpc']}**")
        st.caption(cpc.get(result["cpc"], "CPC sin descripción en el catálogo"))

        target, factor, conversion_message = conversion_for(row, units)
        if target is not None:
            st.info(f"Conversión del catálogo: {source_unit_name} → {units[target]['label']}. Factor: {display_number(factor)}. Las cantidades originales se multiplicarán por este factor.")
            if st.button("Aplicar conversión a todas las cantidades", key=f"convert_{selected}"):
                try:
                    updates = converted_edits(row, units)
                    current.update(updates)
                    st.session_state.mp_revision += 1
                    st.rerun()
                except ValueError as exc:
                    st.error(str(exc))
        elif source_code is not None and units.get(source_code, {}).get("target"):
            st.warning(conversion_message)

        if result["alerts"]:
            st.warning(f"{len(result['alerts'])} alerta(s):\n\n" + "\n\n".join(f"• {message}" for message in result["alerts"]))
        else:
            st.success("Sin alertas calculadas para este registro.")

        st.markdown("#### Comparar y corregir")
        comparison = []
        for original, change, label in QUANTITIES:
            value_field = {"cantidad_inv_ini":"valor_inv_ini", "cantidad_impor":"valor_impor", "cantidad_local":"valor_local", "cantidad_inv_fin":"valor_inv_fin"}.get(original)
            comparison.append({
                "Concepto": label, "Cantidad original": row.get(original),
                "Corrección": current.get(change) or "—",
                "Cantidad efectiva": display_number(result["quantities"][original]),
                "Valor informado (sin edición)": row.get(value_field) if value_field else "—",
                "Precio efectivo": display_number(result["prices"].get(label)),
            })
        st.dataframe(pd.DataFrame(comparison), hide_index=True, width="stretch")

        with st.form(f"edit_{digest}_{selected}_{st.session_state.mp_revision}"):
            left, right = st.columns(2)
            with left:
                new_cpc = st.text_input("Código CPC corregido (13 dígitos; vacío = original)", value=current["cod_cpc_act"])
                unit_choices = [""] + [str(code) for code in sorted(units)]
                chosen_unit = current["uni_med_act"]
                if chosen_unit and chosen_unit not in unit_choices:
                    unit_choices.append(chosen_unit)
                new_unit = st.selectbox(
                    "Unidad corregida (vacía = original)", unit_choices,
                    index=unit_choices.index(chosen_unit),
                    format_func=lambda code: "Conservar unidad original" if not code else f"{code} · {units.get(int(code), {}).get('label', 'Código desconocido')}",
                )
                new_percent = st.text_input("Porcentaje local corregido (vacío = original)", value=current["porcentaje_local_act"])
            with right:
                quantity_edits = {}
                for original, field, label in QUANTITIES:
                    quantity_edits[field] = st.text_input(f"{label}: cantidad corregida", value=current[field], help="Vacío conserva el dato original. Cero es una corrección válida.")
            new_status = st.selectbox("Estado de revisión", REVIEW_STATES, index=REVIEW_STATES.index(current["Estado_revision"]))
            justification = st.text_area("Justificación a las novedades", value=current["Justificación a las novedades"])
            responsible = st.text_input("Responsable", value=current["Responsable"])
            submitted = st.form_submit_button("Guardar revisión", type="primary")
        if submitted:
            updated = {
                "cod_cpc_act": new_cpc.strip(), "uni_med_act": new_unit,
                "porcentaje_local_act": new_percent.strip(),
                "Justificación a las novedades": justification.strip(),
                "Responsable": responsible.strip(), "Estado_revision": new_status,
                **{field: value.strip() for field, value in quantity_edits.items()},
            }
            bad_numbers = [field for field in NUMERIC_EDITS if not blank(updated[field]) and decimal(updated[field]) is None]
            if bad_numbers:
                st.error("Correcciones no numéricas: " + ", ".join(bad_numbers))
            elif updated["cod_cpc_act"] and (len(updated["cod_cpc_act"]) != 13 or not updated["cod_cpc_act"].isdigit()):
                st.error("El CPC corregido debe tener 13 dígitos.")
            else:
                proposed = evaluate(row, updated, units, cpc)
                if new_status == "Validado" and proposed["alerts"] and not justification.strip():
                    st.error("Para validar un registro con alertas, escribe una justificación.")
                else:
                    edits[selected] = updated
                    st.session_state.mp_revision += 1
                    st.rerun()

with tab_company:
    st.subheader("Matriz de control por empresa")
    company_text = st.text_input("Buscar empresa por identificador o razón social", key="company_search").strip().casefold()
    company_alerts = st.checkbox("Solo empresas con alertas", key="company_alerts")
    master_records = company_summary(rows, edits, report)
    master_records = [item for item in master_records
                      if (company == "Todas" or item["Empresa"] == company)
                      and (not company_alerts or item["Con alertas"] > 0)
                      and (not company_text or company_text in (item["Empresa"] + " " + item["Razón social"]).casefold())]
    if not master_records:
        st.info("No hay empresas que cumplan estos filtros.")
    else:
        master = pd.DataFrame(master_records)
        st.caption("Selecciona una fila para ver todas las materias primas de esa empresa. Los filtros de registro de la barra lateral se aplican a la pestaña individual.")
        event = st.dataframe(master, hide_index=True, width="stretch",
                             on_select="rerun", selection_mode="single-row", key=f"master_{digest}")
        indices = event.selection.rows
        company_key = master.iloc[indices[0]]["Empresa"] if indices and indices[0] < len(master) else master.iloc[0]["Empresa"]
        company_keys = [key for key, row in rows.items() if str(row.get("inec_identificador_empresa")) == company_key]
        first = rows[company_keys[0]]
        st.subheader(f"Materias primas de la empresa {company_key}")
        st.write(f"**Razón social:** {first.get('razon_social') or first.get('nombre_comercial') or 'No disponible'} · **RUC:** {first.get('ruc') or 'No disponible'}")
        st.caption(f"{len(company_keys)} registros. Los campos originales y monetarios son de consulta; puedes editar correcciones, estado, justificación y responsable.")
        complete = st.checkbox("Mostrar todas las columnas originales", key=f"complete_{company_key}")
        matrix = pd.DataFrame(matrix_records(company_keys, rows, edits, report, cpc)).set_index("id_mp", drop=False)
        primary_columns = ["id_mp", "nombre", "cpc", "cod_cpc_act", "CPC efectivo", "Descripción CPC efectiva",
                           "unid_med", "uni_med_act", "Unidad efectiva", "Convertir según UM"]
        for original, change, _ in QUANTITIES:
            primary_columns += [original, change]
            money = {"cantidad_inv_ini":"valor_inv_ini", "cantidad_impor":"valor_impor", "cantidad_local":"valor_local", "cantidad_inv_fin":"valor_inv_fin"}.get(original)
            if money:
                primary_columns.append(money)
        primary_columns += ["porcentaje_local", "porcentaje_local_act", "Alertas app", "Estado_revision", "Justificación a las novedades", "Responsable"]
        visible = list(matrix.columns) if complete else primary_columns
        editable_fields = set(EDIT_FIELDS) | {"Convertir según UM"}
        config = {"Estado_revision": st.column_config.SelectboxColumn("Estado revisión", options=list(REVIEW_STATES), required=True),
                  "uni_med_act": st.column_config.SelectboxColumn("Unidad corregida", options=[""] + [str(code) for code in sorted(units)]),
                  "Convertir según UM": st.column_config.CheckboxColumn("Convertir según UM", help="Al guardar, sustituye unidad y cantidades corregidas por la conversión de los datos originales usando UM."),
                  "Alertas app": st.column_config.TextColumn("Alertas calculadas", width="large")}
        with st.expander("Conversiones disponibles para esta empresa"):
            previews = []
            for key in company_keys:
                target, factor, message = conversion_for(rows[key], units)
                if target is not None:
                    previews.append({"id_mp": key, "Origen": str(rows[key].get("unid_med")),
                                     "Destino": units[target]["label"], "Factor": display_number(factor)})
            if previews:
                st.dataframe(pd.DataFrame(previews), hide_index=True, width="stretch")
            else:
                st.info("No hay conversiones con destino y factor definidos en UM para esta empresa.")
        st.caption("Vacío en una corrección conserva el original; 0 es una corrección. Marca Convertir según UM para calcular las cinco cantidades al guardar. Las alertas se actualizan después del guardado.")
        with st.form(f"company_form_{digest}_{company_key}_{st.session_state.mp_revision}_{complete}"):
            edited_matrix = st.data_editor(matrix[visible], hide_index=True, width="stretch",
                                          num_rows="fixed", column_config=config,
                                          disabled=[col for col in visible if col not in editable_fields] + ["_index"],
                                          key=f"company_editor_{digest}_{company_key}_{st.session_state.mp_revision}_{complete}")
            save_company = st.form_submit_button("Guardar cambios de la empresa", type="primary")
        if save_company:
            try:
                pending = prepare_batch(edited_matrix.to_dict("records"), company_keys, rows, edits, units, cpc)
            except ValueError as exc:
                st.error(str(exc))
            else:
                edits.update(pending)
                st.session_state.mp_revision += 1
                st.session_state.mp_message = f"Se guardaron los {len(pending)} registros de la empresa {company_key}."
                st.rerun()

if st.session_state.get("mp_message"):
    st.success(st.session_state.pop("mp_message"))

with st.expander("Resumen general y descarga"):
    summary = []
    for key, row in rows.items():
        result = report[key]
        summary.append({
            "id_mp": key, "Empresa": str(row.get("inec_identificador_empresa") or ""),
            "Materia prima": row.get("nombre"), "CPC efectivo": result["cpc"],
            "Estado": edits[key]["Estado_revision"], "Correcciones": len(result["changes"]),
            "Alertas": len(result["alerts"]), "Detalle": "; ".join(result["alerts"]),
        })
    st.dataframe(pd.DataFrame(summary), hide_index=True, width="stretch")
    st.caption("La descarga conserva la matriz, los catálogos, las correcciones y los campos de revisión. La hoja Revision_app contiene una instantánea de las alertas calculadas.")
    export_key = (digest, st.session_state.mp_revision)
    if st.session_state.get("mp_export_key") != export_key:
        st.session_state.mp_export_bytes = export_progress(st.session_state.mp_content, rows, edits, units, cpc, report)
        st.session_state.mp_export_key = export_key
    output = st.session_state.mp_export_bytes
    st.download_button(
        "Descargar avance (.xlsx)", data=output,
        file_name="ENESEM_Materias_Primas_avance.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
